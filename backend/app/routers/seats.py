from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form, Query, Request
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app import schemas, models, auth, crypto
from app.config import get_settings
from app.audit import log_action
from app.ratelimit import _get_client_ip, file_upload_rate_limiter
import json
import io
import zipfile
import base64
import urllib.parse
import os
import glob
import hmac
import hashlib
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter
import re
import random
import secrets
import time as _time
from contextlib import contextmanager

settings = get_settings()

router = APIRouter(prefix="/api/seats", tags=["seats"])


class _SeatAssignLock:
    _LOCK_TIMEOUT = 30

    def __init__(self):
        self._lock_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
            "data", ".seat_assign.lock"
        )
        self._lock_fd = None

    def acquire(self) -> bool:
        os.makedirs(os.path.dirname(self._lock_path), exist_ok=True)
        deadline = _time.monotonic() + self._LOCK_TIMEOUT
        while _time.monotonic() < deadline:
            try:
                self._lock_fd = open(self._lock_path, 'w')
                if os.name == 'nt':
                    import msvcrt
                    msvcrt.locking(self._lock_fd.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(self._lock_fd.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                return True
            except (IOError, OSError):
                if self._lock_fd:
                    self._lock_fd.close()
                    self._lock_fd = None
                _time.sleep(0.1 + random.random() * 0.1)
        return False

    def release(self):
        if self._lock_fd is None:
            return
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(self._lock_fd.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(self._lock_fd.fileno(), fcntl.LOCK_UN)
        except Exception:
            pass
        finally:
            try:
                self._lock_fd.close()
            except Exception:
                pass
            self._lock_fd = None


@contextmanager
def seat_assign_lock():
    lock = _SeatAssignLock()
    acquired = lock.acquire()
    if not acquired:
        raise HTTPException(
            status_code=status.HTTP_423_LOCKED,
            detail="座位分配操作正在进行中，请稍后重试"
        )
    try:
        yield
    finally:
        lock.release()


def _verify_teacher_password_for_action(db: Session, password: str) -> bool:
    return auth.verify_teacher_password(db, password)


def _get_signing_key() -> bytes:
    raw_key = settings.ENCRYPTION_KEY.encode("utf-8") if isinstance(settings.ENCRYPTION_KEY, str) else settings.ENCRYPTION_KEY
    if len(raw_key) != 32:
        raw_key = hashlib.sha256(raw_key).digest()
    return hmac.new(raw_key, b"signing-key-v5", hashlib.sha256).digest()


def _sign_export(data: str) -> str:
    signing_key = _get_signing_key()
    return hmac.new(signing_key, data.encode("utf-8"), hashlib.sha256).hexdigest()[:32]


def _verify_export_signature(data: str, signature: str) -> bool:
    expected = _sign_export(data)
    return hmac.compare_digest(expected, signature)


def get_seat_state(db: Session) -> models.SeatState:
    state = db.query(models.SeatState).first()
    if not state:
        state = models.SeatState()
        db.add(state)
        db.commit()
        db.refresh(state)
    return state


def seat_to_dict(seat: models.Seat) -> dict:
    return {
        "id": seat.id,
        "label": seat.label,
        "type": seat.seat_type,
        "groupIdx": seat.group_idx,
        "row": seat.row,
        "col": seat.col,
        "isOccupied": seat.is_occupied,
        "isDisabled": seat.is_disabled,
        "occupantName": seat.occupant_name,
        "occupantId": seat.occupant_id,
    }


def student_to_dict(student: models.Student, role: str = "teacher") -> dict:
    if role == "student":
        return {
            "id": student.id,
            "name": student.name,
            "studentId": student.student_id,
            "totalScore": student.total_score,
            "assigned": student.assigned,
            "seatId": student.seat_id,
            "gender": student.gender,
            "selectionMethod": student.selection_method,
            "roundId": student.round_id,
        }
    return {
        "id": student.id,
        "name": student.name,
        "studentId": student.student_id,
        "totalScore": student.total_score,
        "assigned": student.assigned,
        "seatId": student.seat_id,
        "gender": student.gender,
        "selectionMethod": student.selection_method,
        "roundId": student.round_id,
    }


def student_to_export_dict(student: models.Student) -> dict:
    return {
        "id": student.id,
        "name": student.name,
        "studentId": student.student_id,
        "totalScore": student.total_score,
        "assigned": student.assigned,
        "seatId": student.seat_id,
        "gender": student.gender,
        "selectionMethod": student.selection_method,
        "roundId": student.round_id,
    }


def wish_to_dict(wish: models.Wish) -> dict:
    return {
        "id": wish.id,
        "studentId": wish.student_id,
        "studentName": wish.student_name,
        "seat_id": wish.seat_id,
        "priority": wish.priority,
    }


def invitation_to_dict(inv: models.Invitation) -> dict:
    return {
        "id": inv.id,
        "fromStudentId": inv.from_student_id,
        "toStudentId": inv.to_student_id,
        "fromStudentName": inv.from_student_name,
        "toStudentName": inv.to_student_name,
        "status": inv.status,
    }


@router.get("/state")
def get_full_state(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher", "student"))
):
    state = get_seat_state(db)
    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    seats = db.query(models.Seat).order_by(models.Seat.id).all()

    role = current_user["role"]
    sorted_students = sorted(students, key=lambda s: s.total_score, reverse=True)

    if role == "student":
        own_id = current_user.get("student_id")
        own_student = next((s for s in sorted_students if s.id == own_id), None)
        return {
            "seats": [seat_to_dict(s) for s in seats],
            "students": [student_to_dict(own_student, role="student")] if own_student else [],
            "sortedStudentIds": [own_student.id] if own_student else [],
            "currentIndex": state.current_index,
            "globalShowNames": state.global_show_names,
            "companionQuota": state.companion_quota,
            "genderDisplay": state.gender_display,
            "forbiddenRules": state.forbidden_rules or [],
            "mirrorH": state.mirror_h,
            "mirrorV": state.mirror_v,
            "studentListCollapsed": state.student_list_collapsed,
            "completedCollapsed": state.completed_collapsed,
            "groupsCount": state.groups_count,
            "rowsPerGroup": state.rows_per_group,
            "colsPerGroup": state.cols_per_group,
            "hasPairs": state.has_pairs,
        }

    return {
        "seats": [seat_to_dict(s) for s in seats],
        "students": [student_to_dict(s) for s in sorted_students],
        "sortedStudentIds": [s.id for s in sorted_students],
        "currentIndex": state.current_index,
        "globalShowNames": state.global_show_names,
        "companionQuota": state.companion_quota,
        "genderDisplay": state.gender_display,
        "forbiddenRules": state.forbidden_rules or [],
        "mirrorH": state.mirror_h,
        "mirrorV": state.mirror_v,
        "studentListCollapsed": state.student_list_collapsed,
        "completedCollapsed": state.completed_collapsed,
        "groupsCount": state.groups_count,
        "rowsPerGroup": state.rows_per_group,
        "colsPerGroup": state.cols_per_group,
        "hasPairs": state.has_pairs,
    }


@router.get("/students")
def get_students(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher", "student"))
):
    role = current_user["role"]
    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    if role == "student":
        own_id = current_user.get("student_id")
        own_student = next((s for s in students if s.id == own_id), None)
        return {"students": [student_to_dict(own_student, role="student")] if own_student else []}
    return {"students": [student_to_dict(s) for s in students]}


@router.get("/students/names")
def get_student_names(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher", "student"))
):
    role = current_user["role"]
    students = db.query(models.Student).order_by(models.Student.name).all()
    if role == "student":
        own_name = current_user.get("student_name", "")
        return {"names": [own_name]}
    return {"names": [s.name for s in students]}


@router.get("/students/status")
def get_all_students_with_status(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    students = db.query(models.Student).order_by(models.Student.rank, models.Student.name).all()
    result = []
    for s in students:
        wishes = db.query(models.Wish).filter(models.Wish.student_name == s.name).count()
        result.append({
            "id": s.id,
            "name": s.name,
            "studentId": s.student_id,
            "rank": s.rank,
            "hasPassword": s.password_hash is not None,
            "wishCount": wishes,
            "assigned": s.assigned,
        })
    return {
        "students": result,
        "totalStudents": len(students),
        "hasPassword": sum(1 for r in result if r["hasPassword"]),
        "noPassword": sum(1 for r in result if not r["hasPassword"]),
    }


@router.post("/students/reset-password")
def reset_student_password(
    req: schemas.ResetPasswordRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    student = db.query(models.Student).filter(models.Student.name == req.student_name).first()
    if not student:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="学生不存在")

    student.password_hash = None
    student.password_updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"success": True, "message": f"已清除{req.student_name}的密码，学生下次登录时需重新设置"}


@router.get("/wishes")
def get_all_wishes(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    wishes = db.query(models.Wish).all()
    return {"wishes": [wish_to_dict(w) for w in wishes]}


@router.get("/wishes/{student_name}")
def get_student_wishes(
    student_name: str,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("student", "teacher"))
):
    if current_user["role"] == "student":
        own_name = current_user.get("student_name")
        if own_name and own_name != student_name:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="只能查看自己的志愿")
    wishes = db.query(models.Wish).filter(models.Wish.student_name == student_name).order_by(models.Wish.priority).all()
    return {"wishes": [wish_to_dict(w) for w in wishes]}


@router.post("/wishes/save")
def save_student_wishes(
    req: schemas.SaveWishesRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("student"))
):
    student_name = current_user.get("student_name")
    student_id = current_user.get("student_id")
    if not student_name or not student_id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="无法识别学生身份")

    active_round = _get_active_round(db)
    if not active_round:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="当前不在选座周期中，无法设置志愿"
        )

    db.query(models.Wish).filter(models.Wish.student_id == student_id).delete()

    for w in req.wishes:
        wish = models.Wish(
            student_id=student_id,
            student_name=student_name,
            seat_id=w.seat_id,
            priority=w.priority,
        )
        db.add(wish)

    db.commit()
    return {"success": True, "message": "志愿保存成功", "count": len(req.wishes)}


@router.post("/assign")
def assign_seat(
    req: schemas.SeatAssignRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    with seat_assign_lock():
        return _do_assign_seat(req, db, current_user)


def _do_assign_seat(req: schemas.SeatAssignRequest, db: Session, current_user: dict):
    state = get_seat_state(db)
    seat = db.query(models.Seat).filter(models.Seat.id == req.seat_id).first()
    if not seat:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="座位不存在")
    if seat.is_disabled:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="该座位已禁用")

    student_name = req.student_name
    if not student_name:
        students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
        sorted_list = sorted(students, key=lambda s: s.total_score, reverse=True)
        idx = state.current_index
        if idx >= len(sorted_list):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="所有学生已分配完成")
        student = sorted_list[idx]
        student_name = student.name
    else:
        student = db.query(models.Student).filter(models.Student.name == student_name).first()
        if not student:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="学生不存在")

    student = db.query(models.Student).filter(models.Student.name == student_name).first()
    if student.assigned:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="该学生已分配座位")

    if seat.is_occupied:
        old_student = db.query(models.Student).filter(models.Student.name == seat.occupant_name).first()
        if old_student:
            old_student.assigned = False
            old_student.seat_id = None

    seat.is_occupied = True
    seat.occupant_name = student_name
    seat.occupant_id = student.id
    student.assigned = True
    student.seat_id = seat.id

    active_round = db.query(models.SelectionRound).filter(
        models.SelectionRound.status == "active"
    ).first()
    if active_round:
        student.round_id = active_round.id
        wish = db.query(models.Wish).filter(
            models.Wish.student_id == student.id,
            models.Wish.seat_id == seat.id
        ).first()
        if wish:
            student.selection_method = f"第{wish.priority}志愿"
        else:
            student.selection_method = "当场选座"

    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    sorted_list = sorted(students, key=lambda s: s.total_score, reverse=True)
    while state.current_index < len(sorted_list) and sorted_list[state.current_index].assigned:
        state.current_index += 1

    if req.companion_name:
        companion = db.query(models.Student).filter(models.Student.name == req.companion_name).first()
        if companion and not companion.assigned:
            adjacent = _find_adjacent_seat(db, seat)
            if adjacent:
                adjacent.is_occupied = True
                adjacent.occupant_name = req.companion_name
                adjacent.occupant_id = companion.id
                companion.assigned = True
                companion.seat_id = adjacent.id
                if active_round:
                    companion.round_id = active_round.id
                    companion_wish = db.query(models.Wish).filter(
                        models.Wish.student_id == companion.id,
                        models.Wish.seat_id == adjacent.id
                    ).first()
                    if companion_wish:
                        companion.selection_method = f"第{companion_wish.priority}志愿"
                    else:
                        companion.selection_method = "当场选座"

    db.commit()
    return {"success": True, "message": f"{student_name} 已分配到 {seat.label}"}


def _find_adjacent_seat(db: Session, seat: models.Seat) -> models.Seat | None:
    if seat.seat_type != "group" or seat.row is None or seat.col is None or seat.group_idx is None:
        return None
    candidates = db.query(models.Seat).filter(
        models.Seat.seat_type == "group",
        models.Seat.group_idx == seat.group_idx,
        models.Seat.row == seat.row,
        models.Seat.col.in_([seat.col - 1, seat.col + 1]),
        models.Seat.is_occupied == False,
        models.Seat.is_disabled == False,
    ).all()
    return candidates[0] if candidates else None


@router.post("/swap")
def swap_seats(
    req: schemas.SeatSwapRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    src = db.query(models.Seat).filter(models.Seat.id == req.src_seat_id).first()
    dst = db.query(models.Seat).filter(models.Seat.id == req.dst_seat_id).first()
    if not src or not dst:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="座位不存在")
    if not src.is_occupied:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="源座位无学生")

    if dst.is_disabled and not dst.is_occupied:
        dst.occupant_name = src.occupant_name
        dst.occupant_id = src.occupant_id
        dst.is_occupied = True
        dst.is_disabled = False
        src.occupant_name = None
        src.occupant_id = None
        src.is_occupied = False
        src.is_disabled = True

        student = db.query(models.Student).filter(models.Student.name == dst.occupant_name).first()
        if student:
            student.seat_id = dst.id
    elif dst.is_occupied:
        src_name = src.occupant_name
        dst_name = dst.occupant_name
        src_id = src.occupant_id
        dst_id = dst.occupant_id

        src.occupant_name = dst_name
        src.occupant_id = dst_id
        dst.occupant_name = src_name
        dst.occupant_id = src_id

        s1 = db.query(models.Student).filter(models.Student.name == src_name).first()
        s2 = db.query(models.Student).filter(models.Student.name == dst_name).first()
        if s1:
            s1.seat_id = dst.id
        if s2:
            s2.seat_id = src.id
    else:
        dst.occupant_name = src.occupant_name
        dst.occupant_id = src.occupant_id
        dst.is_occupied = True
        src.occupant_name = None
        src.occupant_id = None
        src.is_occupied = False

        student = db.query(models.Student).filter(models.Student.name == dst.occupant_name).first()
        if student:
            student.seat_id = dst.id

    db.commit()
    return {"success": True, "message": "座位交换完成"}


@router.post("/clear")
def clear_seat(
    req: schemas.SeatClearRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    seat = db.query(models.Seat).filter(models.Seat.id == req.seat_id).first()
    if not seat:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="座位不存在")
    if not seat.is_occupied:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="座位上没有学生")

    student = db.query(models.Student).filter(models.Student.name == seat.occupant_name).first()
    if student:
        student.assigned = False
        student.seat_id = None
        student.selection_method = None
        student.round_id = None

    seat.is_occupied = False
    seat.occupant_name = None
    seat.occupant_id = None

    state = get_seat_state(db)
    sorted_list = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    sorted_list = sorted(sorted_list, key=lambda s: s.total_score, reverse=True)
    for i in range(len(sorted_list)):
        if not sorted_list[i].assigned and i < state.current_index:
            state.current_index = i
            break

    db.commit()
    return {"success": True, "message": "已清除座位"}


@router.post("/next")
def next_student(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    state = get_seat_state(db)
    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    sorted_list = sorted(students, key=lambda s: s.total_score, reverse=True)

    while state.current_index < len(sorted_list) and sorted_list[state.current_index].assigned:
        state.current_index += 1

    if state.current_index >= len(sorted_list):
        db.commit()
        return {"success": True, "finished": True, "message": "所有学生已完成选座"}

    current = sorted_list[state.current_index]
    db.commit()
    return {
        "success": True,
        "finished": False,
        "currentStudent": student_to_dict(current),
        "currentIndex": state.current_index,
    }


@router.post("/undo")
def undo_assignment(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    state = get_seat_state(db)
    if state.current_index <= 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="没有可撤销的操作")

    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    sorted_list = sorted(students, key=lambda s: s.total_score, reverse=True)

    prev_idx = state.current_index - 1
    if prev_idx >= 0 and prev_idx < len(sorted_list):
        prev_student = sorted_list[prev_idx]
        if prev_student.seat_id:
            seat = db.query(models.Seat).filter(models.Seat.id == prev_student.seat_id).first()
            if seat:
                seat.is_occupied = False
                seat.occupant_name = None
                seat.occupant_id = None
            prev_student.assigned = False
            prev_student.seat_id = None
            prev_student.selection_method = None
            prev_student.round_id = None
        state.current_index = prev_idx

    db.commit()
    return {"success": True, "message": "已撤销上一步", "currentIndex": state.current_index}


@router.put("/state")
def update_seat_state(
    req: schemas.SeatStateUpdate,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    state = get_seat_state(db)

    update_data = req.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        if hasattr(state, key):
            setattr(state, key, value)

    db.commit()
    return {"success": True, "message": "状态已更新"}


@router.post("/reset")
def reset_all(
    req: schemas.DestructiveActionConfirm,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if not auth.verify_teacher_password(db, req.teacher_password):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="教师密码验证失败")

    state = get_seat_state(db)
    state.current_index = 0

    seats = db.query(models.Seat).all()
    for seat in seats:
        seat.is_occupied = False
        seat.occupant_name = None
        seat.occupant_id = None

    students = db.query(models.Student).all()
    for student in students:
        student.assigned = False
        student.seat_id = None
        student.selection_method = None
        student.round_id = None

    db.commit()

    log_action(
        action="reset_seats",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"students_count": len(students), "seats_count": len(seats)},
        ip=_get_client_ip(request) if request else None,
    )

    return {"success": True, "message": "所有选座数据已重置"}


def _get_data_dir() -> str:
    data_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "data")
    os.makedirs(data_dir, exist_ok=True)
    return data_dir


def _auto_export_backup(db: Session) -> str:
    data_dir = _get_data_dir()
    json_str = _build_export_json(db)
    encrypted = crypto.encrypt_data(json_str)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    rand = secrets.token_hex(8)
    filename = f"auto_backup_{timestamp}_{rand}.json"
    filepath = os.path.join(data_dir, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(encrypted)
    return filepath


def _get_latest_backup() -> str | None:
    data_dir = _get_data_dir()
    files = glob.glob(os.path.join(data_dir, "auto_backup_*.json"))
    if not files:
        return None
    return max(files, key=os.path.getmtime)


def _strip_signature(json_str: str) -> str:
    try:
        obj = json.loads(json_str)
        obj.pop("_signature", None)
        return json.dumps(obj, ensure_ascii=False, sort_keys=True)
    except Exception:
        return json_str


def _backup_matches_current(db: Session, backup_path: str) -> bool:
    if not os.path.exists(backup_path):
        return False
    current_json = _build_export_json(db, include_signature=False)
    with open(backup_path, "r", encoding="utf-8") as f:
        encrypted_backup = f.read()
    decrypted = crypto.decrypt_data(encrypted_backup)
    if decrypted is None:
        return False
    return current_json == _strip_signature(decrypted)


def _get_active_round(db: Session) -> models.SelectionRound | None:
    return db.query(models.SelectionRound).filter(
        models.SelectionRound.status == "active"
    ).first()


def _get_next_round_number(db: Session) -> int:
    last = db.query(models.SelectionRound).order_by(
        models.SelectionRound.round_number.desc()
    ).first()
    return (last.round_number + 1) if last else 1


@router.get("/rounds")
def list_rounds(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher", "student"))
):
    role = current_user["role"]
    rounds = db.query(models.SelectionRound).order_by(
        models.SelectionRound.round_number.desc()
    ).all()
    if role == "student":
        active = _get_active_round(db)
        if not active:
            return {"rounds": []}
        return {
            "rounds": [
                {
                    "id": active.id,
                    "roundNumber": active.round_number,
                    "name": active.name,
                    "status": active.status,
                    "noWishLockDate": active.no_wish_lock_date,
                    "classChangeDate": active.class_change_date,
                    "createdAt": active.created_at.isoformat() if active.created_at else None,
                    "completedAt": active.completed_at.isoformat() if active.completed_at else None,
                }
            ]
        }
    return {
        "rounds": [
            {
                "id": r.id,
                "roundNumber": r.round_number,
                "name": r.name,
                "status": r.status,
                "noWishLockDate": r.no_wish_lock_date,
                "classChangeDate": r.class_change_date,
                "createdAt": r.created_at.isoformat() if r.created_at else None,
                "completedAt": r.completed_at.isoformat() if r.completed_at else None,
            }
            for r in rounds
        ]
    }


@router.get("/rounds/current")
def get_current_round(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher", "student"))
):
    active = _get_active_round(db)
    if not active:
        return {"round": None}
    return {
        "round": {
            "id": active.id,
            "roundNumber": active.round_number,
            "name": active.name,
            "status": active.status,
            "noWishLockDate": active.no_wish_lock_date,
            "classChangeDate": active.class_change_date,
            "createdAt": active.created_at.isoformat() if active.created_at else None,
            "completedAt": active.completed_at.isoformat() if active.completed_at else None,
        }
    }


@router.post("/rounds/start")
def start_new_round(
    req: schemas.RoundCreateRequest = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if req is None:
        req = schemas.RoundCreateRequest()

    existing_active = _get_active_round(db)
    if existing_active:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="当前有进行中的轮次，请先完成当前轮次"
        )

    latest_backup = _get_latest_backup()
    if latest_backup and not _backup_matches_current(db, latest_backup):
        _auto_export_backup(db)

    state = get_seat_state(db)
    state.current_index = 0

    seats = db.query(models.Seat).all()
    for seat in seats:
        seat.is_occupied = False
        seat.occupant_name = None
        seat.occupant_id = None

    students = db.query(models.Student).all()
    for student in students:
        student.assigned = False
        student.seat_id = None
        student.selection_method = None
        student.round_id = None

    db.query(models.Wish).delete()
    db.query(models.Invitation).delete()

    round_number = _get_next_round_number(db)
    name = req.name or f"第{round_number}轮选座"

    new_round = models.SelectionRound(
        round_number=round_number,
        name=name,
        status="active",
        no_wish_lock_date=req.no_wish_lock_date,
        class_change_date=req.class_change_date,
    )
    db.add(new_round)
    db.commit()
    db.refresh(new_round)

    return {
        "success": True,
        "message": f"已开启 {name}",
        "round": {
            "id": new_round.id,
            "roundNumber": new_round.round_number,
            "name": new_round.name,
            "status": new_round.status,
            "noWishLockDate": new_round.no_wish_lock_date,
            "classChangeDate": new_round.class_change_date,
        }
    }


@router.post("/rounds/{round_id}/complete")
def complete_round(
    round_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    round_obj = db.query(models.SelectionRound).filter(
        models.SelectionRound.id == round_id
    ).first()
    if not round_obj:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="轮次不存在")
    if round_obj.status != "active":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="该轮次不是进行中状态")

    round_obj.status = "completed"
    round_obj.completed_at = datetime.now(timezone.utc)

    backup_path = _auto_export_backup(db)
    round_obj.backup_path = backup_path

    db.commit()

    return {
        "success": True,
        "message": f"{round_obj.name} 已完成",
        "backupPath": backup_path,
    }


@router.put("/rounds/{round_id}")
def update_round(
    round_id: int,
    req: schemas.RoundUpdateRequest,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    round_obj = db.query(models.SelectionRound).filter(
        models.SelectionRound.id == round_id
    ).first()
    if not round_obj:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="轮次不存在")

    if req.name is not None:
        round_obj.name = req.name
    if req.no_wish_lock_date is not None:
        round_obj.no_wish_lock_date = req.no_wish_lock_date
    if req.class_change_date is not None:
        round_obj.class_change_date = req.class_change_date

    db.commit()
    return {"success": True, "message": "轮次信息已更新"}


@router.get("/rounds/{round_id}/results")
def get_round_results(
    round_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    round_obj = db.query(models.SelectionRound).filter(
        models.SelectionRound.id == round_id
    ).first()
    if not round_obj:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="轮次不存在")

    if round_obj.status == "active":
        state = get_seat_state(db)
        students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
        seats = db.query(models.Seat).order_by(models.Seat.id).all()
        sorted_students = sorted(students, key=lambda s: s.total_score, reverse=True)
        return {
            "round": {
                "id": round_obj.id,
                "roundNumber": round_obj.round_number,
                "name": round_obj.name,
                "status": round_obj.status,
                "noWishLockDate": round_obj.no_wish_lock_date,
                "classChangeDate": round_obj.class_change_date,
            },
            "seats": [seat_to_dict(s) for s in seats],
            "students": [student_to_dict(s) for s in sorted_students],
            "groupsCount": state.groups_count,
            "rowsPerGroup": state.rows_per_group,
            "colsPerGroup": state.cols_per_group,
            "hasPairs": state.has_pairs,
        }

    if not round_obj.backup_path or not os.path.exists(round_obj.backup_path):
        return {
            "round": {
                "id": round_obj.id,
                "roundNumber": round_obj.round_number,
                "name": round_obj.name,
                "status": round_obj.status,
                "noWishLockDate": round_obj.no_wish_lock_date,
                "classChangeDate": round_obj.class_change_date,
            },
            "seats": [],
            "students": [],
            "groupsCount": 4,
            "rowsPerGroup": 7,
            "colsPerGroup": 2,
            "hasPairs": True,
        }

    with open(round_obj.backup_path, "r", encoding="utf-8") as f:
        encrypted = f.read()
    decrypted = crypto.decrypt_data(encrypted)
    if not decrypted:
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="无法解密轮次数据")

    data = json.loads(decrypted)
    return {
        "round": {
            "id": round_obj.id,
            "roundNumber": round_obj.round_number,
            "name": round_obj.name,
            "status": round_obj.status,
            "noWishLockDate": round_obj.no_wish_lock_date,
            "classChangeDate": round_obj.class_change_date,
        },
        "seats": data.get("seats", []),
        "students": data.get("students", []),
        "groupsCount": data.get("groupsCount", 4),
        "rowsPerGroup": data.get("rowsPerGroup", 7),
        "colsPerGroup": data.get("colsPerGroup", 2),
        "hasPairs": data.get("hasPairs", True),
    }


def _build_export_json(db: Session, include_signature: bool = True) -> str:
    state = get_seat_state(db)
    students = db.query(models.Student).order_by(models.Student.total_score.desc()).all()
    seats = db.query(models.Seat).order_by(models.Seat.id).all()
    wishes = db.query(models.Wish).all()
    rounds = db.query(models.SelectionRound).order_by(models.SelectionRound.round_number).all()

    export_obj = {
        "seats": [seat_to_dict(s) for s in seats],
        "students": [student_to_export_dict(s) for s in students],
        "sortedStudentIds": [s.id for s in sorted(students, key=lambda x: x.total_score, reverse=True)],
        "currentIndex": state.current_index,
        "companionQuota": state.companion_quota,
        "genderDisplay": state.gender_display,
        "forbiddenRules": state.forbidden_rules,
        "globalShowNames": state.global_show_names,
        "mirrorH": state.mirror_h,
        "mirrorV": state.mirror_v,
        "studentListCollapsed": state.student_list_collapsed,
        "completedCollapsed": state.completed_collapsed,
        "groupsCount": state.groups_count,
        "rowsPerGroup": state.rows_per_group,
        "colsPerGroup": state.cols_per_group,
        "hasPairs": state.has_pairs,
        "studentWishes": {
            w.student_name: [
                {"seat_id": w2.seat_id, "priority": w2.priority}
                for w2 in wishes if w2.student_name == w.student_name
            ]
            for w in wishes
        },
        "rounds": [
            {
                "id": r.id,
                "roundNumber": r.round_number,
                "name": r.name,
                "status": r.status,
                "noWishLockDate": r.no_wish_lock_date,
                "classChangeDate": r.class_change_date,
                "createdAt": r.created_at.isoformat() if r.created_at else None,
                "completedAt": r.completed_at.isoformat() if r.completed_at else None,
            }
            for r in rounds
        ],
        "_version": 303,
    }

    if include_signature:
        data_json = json.dumps(export_obj, ensure_ascii=False, sort_keys=True)
        signature = _sign_export(data_json)
        export_obj["_signature"] = signature

    return json.dumps(export_obj, ensure_ascii=False, sort_keys=True)


def _create_csc123_file(json_str: str, password: str | None) -> bytes:
    buf = io.BytesIO()
    zip_mode = zipfile.ZIP_DEFLATED

    if password:
        inner_data = crypto.encrypt_with_password(json_str, password)
    else:
        inner_data = json_str

    with zipfile.ZipFile(buf, 'w', zip_mode) as zf:
        if password:
            zf.setpassword(password.encode("utf-8"))
        zf.writestr("data.json", inner_data)

    return buf.getvalue()


_MAX_DECOMPRESSED_SIZE = 50 * 1024 * 1024


def _check_zip_bomb(zf: zipfile.ZipFile) -> tuple[bool, int]:
    total_size = 0
    for info in zf.infolist():
        total_size += info.file_size
        if total_size > _MAX_DECOMPRESSED_SIZE:
            return False, total_size
    return True, total_size


def _extract_csc123_file(file_bytes: bytes, password: str | None) -> tuple[str | None, bool]:
    try:
        buf = io.BytesIO(file_bytes)
        with zipfile.ZipFile(buf, 'r') as zf:
            ok, total = _check_zip_bomb(zf)
            if not ok:
                return None, False
            namelist = zf.namelist()
            if "data.json" not in namelist:
                return None, False
            if password:
                zf.setpassword(password.encode("utf-8"))
            raw = zf.read("data.json")
    except RuntimeError as e:
        if "password" in str(e).lower():
            return None, True
        return None, False
    except Exception:
        return None, False

    try:
        raw_str = raw.decode("utf-8")
    except UnicodeDecodeError:
        return None, False

    if not password and not raw_str.startswith("{"):
        return None, True

    return raw_str, False


@router.get("/export-xlsx")
def export_xlsx(
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    seats = db.query(models.Seat).order_by(models.Seat.id).all()

    groups = {}
    pair_seats = []
    for s in seats:
        if s.seat_type == "group" and s.group_idx is not None:
            groups.setdefault(s.group_idx, []).append(s)
        elif s.seat_type == "pair":
            pair_seats.append(s)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "座位表"

    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    occupied_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    disabled_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    empty_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    group_header_font = Font(bold=True, size=12)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    max_rows = 0
    for g_seats in groups.values():
        if g_seats:
            mr = max((s.row or 0) for s in g_seats) + 1
            if mr > max_rows:
                max_rows = mr

    col_offset = 0
    for g_idx in sorted(groups.keys()):
        g_seats = sorted(groups[g_idx], key=lambda x: (x.row or 0, x.col or 0))
        if not g_seats:
            continue
        max_row = max(s.row or 0 for s in g_seats)
        max_col = max(s.col or 0 for s in g_seats)

        ws.cell(row=1, column=col_offset + 1, value=f"第 {g_idx + 1} 组").font = group_header_font
        ws.merge_cells(start_row=1, start_column=col_offset + 1, end_row=1, end_column=col_offset + max_col + 1)

        for r in range(max_row + 1):
            for c in range(max_col + 1):
                cell = ws.cell(row=r + 2, column=col_offset + c + 1)
                cell.border = thin_border
                cell.alignment = center_align

                seat = next((s for s in g_seats if s.row == r and s.col == c), None)
                if seat:
                    if seat.is_disabled:
                        cell.value = "×"
                        cell.fill = disabled_fill
                    elif seat.is_occupied and seat.occupant_name:
                        cell.value = seat.occupant_name
                        cell.fill = occupied_fill
                    else:
                        cell.fill = empty_fill
                else:
                    cell.fill = empty_fill

        col_offset += max_col + 2

    if pair_seats:
        ws.cell(row=max_rows + 4, column=1, value="独立座位").font = group_header_font
        for i, seat in enumerate(pair_seats):
            cell = ws.cell(row=max_rows + 5, column=i + 1)
            cell.border = thin_border
            cell.alignment = center_align
            if seat.is_disabled:
                cell.value = "×"
                cell.fill = disabled_fill
            elif seat.is_occupied and seat.occupant_name:
                cell.value = seat.occupant_name
                cell.fill = occupied_fill
            else:
                cell.fill = empty_fill

    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"座位表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    encoded_filename = urllib.parse.quote(filename, safe='')

    from fastapi.responses import Response
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


@router.post("/export")
def export_data(
    req: schemas.ExportDataRequest,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    password = req.password
    json_str = _build_export_json(db)
    file_bytes = _create_csc123_file(json_str, password if password else None)

    if password:
        verified = _extract_csc123_file(file_bytes, password)
        if verified[0] is None:
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="导出校验失败")

    encoded = base64.b64encode(file_bytes).decode("ascii")

    filename = f"seatchoose_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csc123"

    log_action(
        action="export_data",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"filename": filename, "encrypted": bool(password)},
        ip=_get_client_ip(request) if request else None,
    )

    return {
        "success": True,
        "filename": filename,
        "data": encoded,
    }


def _validate_file_magic(raw: bytes, expected_magic: bytes, file_type: str) -> None:
    if len(raw) < len(expected_magic):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"文件内容无效，无法识别为{file_type}文件"
        )
    if raw[:len(expected_magic)] != expected_magic:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"文件内容与扩展名不匹配，请上传有效的{file_type}文件"
        )


@router.post("/import")
def import_data(
    file: UploadFile = File(...),
    password: str | None = Form(None),
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    client_ip = _get_client_ip(request) if request else "unknown"
    file_upload_rate_limiter.check(f"upload:{client_ip}")

    if not file.filename or not file.filename.lower().endswith(".csc123"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持 .csc123 文件")

    if file.content_type and not file.content_type.startswith("application/"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="不支持的文件类型")

    file_bytes = file.file.read()
    _validate_file_magic(file_bytes, b"PK\x03\x04", ".csc123")
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(file_bytes) > max_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"文件大小超过限制（最大{settings.MAX_UPLOAD_SIZE_MB}MB）")
    raw_str, needs_password = _extract_csc123_file(file_bytes, password if password else None)

    if needs_password and not password:
        return {"success": False, "needsPassword": True, "message": "文件已加密，请输入密码"}

    if raw_str is None:
        if password:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="密码错误，无法解密")
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件格式不正确")

    if raw_str.startswith("{"):
        import_obj = json.loads(raw_str)
        _apply_import(db, import_obj)
        log_action(
            action="import_data",
            user_role=current_user["role"],
            user_identity=current_user.get("sub", "unknown"),
            detail={"filename": file.filename},
            ip=_get_client_ip(request) if request else None,
        )
        return {"success": True, "message": "数据导入成功"}

    decrypted = crypto.decrypt_with_password(raw_str, password)
    if decrypted is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="密码错误，无法解密")

    import_obj = json.loads(decrypted)
    _apply_import(db, import_obj)
    log_action(
        action="import_data",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"filename": file.filename, "encrypted": True},
        ip=_get_client_ip(request) if request else None,
    )
    return {"success": True, "message": "数据导入成功"}


def _verify_and_strip_signature(obj: dict) -> None:
    signature = obj.pop("_signature", None)
    if signature:
        data_to_verify = json.dumps(obj, ensure_ascii=False, sort_keys=True)
        if not _verify_export_signature(data_to_verify, signature):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="文件签名校验失败，文件可能已被篡改"
            )


def _apply_import(db: Session, obj: dict):
    try:
        _verify_and_strip_signature(obj)
        state = get_seat_state(db)

        if "students" in obj:
            student_ids_in_file = set()
            student_names_in_file = set()
            for s in obj["students"]:
                sid = s.get("id", 0)
                sname = s.get("name", "")
                if sid in student_ids_in_file:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"导入数据中存在重复的学生ID: {sid}"
                    )
                if sname in student_names_in_file:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"导入数据中存在重复的学生姓名: {sname}"
                    )
                student_ids_in_file.add(sid)
                student_names_in_file.add(sname)

            db.query(models.Student).delete()
            db.flush()
            for s in obj["students"]:
                student = models.Student(
                    id=s.get("id", 0),
                    name=s.get("name", ""),
                    student_id=s.get("studentId"),
                    total_score=s.get("totalScore", 0),
                    assigned=s.get("assigned", False),
                    seat_id=s.get("seatId"),
                    gender=s.get("gender"),
                    selection_method=s.get("selectionMethod"),
                    round_id=s.get("roundId"),
                )
                db.add(student)

        if "seats" in obj:
            seat_ids_in_file = set()
            for s in obj["seats"]:
                sid = s.get("id", 0)
                if sid in seat_ids_in_file:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"导入数据中存在重复的座位ID: {sid}"
                    )
                seat_ids_in_file.add(sid)

            db.query(models.Seat).delete()
            db.flush()
            for s in obj["seats"]:
                seat = models.Seat(
                    id=s.get("id", 0),
                    label=s.get("label", ""),
                    seat_type=s.get("type", "group"),
                    group_idx=s.get("groupIdx"),
                    row=s.get("row"),
                    col=s.get("col"),
                    is_occupied=s.get("isOccupied", False),
                    is_disabled=s.get("isDisabled", False),
                    occupant_name=s.get("occupantName"),
                    occupant_id=s.get("occupantId"),
                )
                db.add(seat)

        if "currentIndex" in obj:
            state.current_index = obj["currentIndex"]
        if "companionQuota" in obj:
            state.companion_quota = obj["companionQuota"]
        if "genderDisplay" in obj:
            state.gender_display = obj["genderDisplay"]
        if "forbiddenRules" in obj:
            state.forbidden_rules = obj["forbiddenRules"]
        if "globalShowNames" in obj:
            state.global_show_names = obj["globalShowNames"]
        if "mirrorH" in obj:
            state.mirror_h = obj["mirrorH"]
        if "mirrorV" in obj:
            state.mirror_v = obj["mirrorV"]
        if "studentListCollapsed" in obj:
            state.student_list_collapsed = obj["studentListCollapsed"]
        if "completedCollapsed" in obj:
            state.completed_collapsed = obj["completedCollapsed"]
        if "groupsCount" in obj:
            state.groups_count = obj["groupsCount"]
        if "rowsPerGroup" in obj:
            state.rows_per_group = obj["rowsPerGroup"]
        if "colsPerGroup" in obj:
            state.cols_per_group = obj["colsPerGroup"]
        if "hasPairs" in obj:
            state.has_pairs = obj["hasPairs"]

        if "studentWishes" in obj:
            seat_ids = _get_import_seat_ids(db, obj)
            db.query(models.Wish).delete()
            db.flush()
            for student_name, wish_list in obj["studentWishes"].items():
                student = db.query(models.Student).filter(models.Student.name == student_name).first()
                if not student:
                    continue
                for w in wish_list:
                    wish_seat_id = w.get("seat_id", w.get("seatId", 0))
                    if wish_seat_id and wish_seat_id not in seat_ids:
                        raise HTTPException(
                            status_code=status.HTTP_400_BAD_REQUEST,
                            detail=f"导入数据中学生 {student_name} 的志愿引用了不存在的座位ID: {wish_seat_id}"
                        )
                    wish = models.Wish(
                        student_id=student.id,
                        student_name=student_name,
                        seat_id=wish_seat_id,
                        priority=w.get("priority", 1),
                    )
                    db.add(wish)

        if "students" in obj and "seats" in obj:
            seat_ids_after = {s.id for s in db.query(models.Seat).all()}
            for s in obj["students"]:
                occupant_seat_id = s.get("seatId")
                if occupant_seat_id is not None and occupant_seat_id not in seat_ids_after:
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail=f"导入数据中学生 {s.get('name', '')} 引用了不存在的座位ID: {occupant_seat_id}"
                    )

        db.commit()
    except Exception:
        db.rollback()
        raise


def _get_import_seat_ids(db: Session, obj: dict) -> set:
    if "seats" in obj:
        return {s.get("id", 0) for s in obj["seats"]}
    return {s.id for s in db.query(models.Seat).all()}


HEADER_ALIASES = {
    "name": re.compile(r"姓名|名字|学生|name", re.I),
    "rank": re.compile(r"排名|名次|rank\s*$", re.I),
    "score": re.compile(r"总分|成绩|分数|score\s|^score$", re.I),
    "student_id": re.compile(r"学号|编号|考号|id|no", re.I),
    "gender": re.compile(r"性别|gender", re.I),
}


def _match_header(value):
    if not value:
        return None
    value = str(value).strip()
    for key, pattern in HEADER_ALIASES.items():
        if pattern.search(value):
            return key
    return None


_GENDER_KEYWORDS = [
    ("大屌萌妹", "m"),
    ("跨性别男性", "m"),
    ("跨性别男", "m"),
    ("女装大佬", "m"),
    ("男の娘", "m"),
    ("女装子", "m"),
    ("男妹子", "m"),
    ("男娘", "m"),
    ("伪娘", "m"),
    ("药娘", "m"),
    ("女装", "m"),
    ("跨男", "m"),
    ("FTM", "m"),
    ("ftm", "m"),
    ("跨性别女性", "f"),
    ("跨性别女", "f"),
    ("变性女", "f"),
    ("跨女", "f"),
    ("伪男", "f"),
    ("跨性别", "f"),
    ("变性", "f"),
    ("MTF", "f"),
    ("mtf", "f"),
    ("男", "m"),
    ("女", "f"),
    ("m", "m"),
    ("M", "m"),
    ("f", "f"),
    ("F", "f"),
    ("1", "m"),
    ("0", "f"),
]


def _normalize_gender(gv: str):
    gv = gv.strip()
    for keyword, gender in _GENDER_KEYWORDS:
        if keyword in gv:
            return gender
    return None


MAX_XLSX_ROWS = 500


def _parse_xlsx_students(file_bytes: bytes):
    ok, _ = _check_zip_bomb(zipfile.ZipFile(io.BytesIO(file_bytes), 'r'))
    if not ok:
        return [], "文件解压后体积过大，拒绝处理"

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    rows = []
    for i, row in enumerate(ws.iter_rows(max_col=min(ws.max_column, 50), values_only=True)):
        if i > MAX_XLSX_ROWS:
            break
        rows.append(row)
    wb.close()

    if len(rows) > MAX_XLSX_ROWS:
        return [], f"文件行数超过{MAX_XLSX_ROWS}行限制"

    if not rows:
        return [], "文件为空"

    header_row = rows[0]
    groups = []
    current_group = []
    for col_idx, cell in enumerate(header_row):
        cell_type = _match_header(cell)
        if cell_type == "name":
            if current_group:
                groups.append(current_group)
            current_group = [("name", col_idx)]
        elif cell_type and current_group:
            if cell_type not in [m[0] for m in current_group]:
                current_group.append((cell_type, col_idx))
    if current_group and any(m[0] == "name" for m in current_group):
        groups.append(current_group)

    if not groups:
        return [], "未识别到姓名列，请确保第一行有'姓名'标题"

    students = []
    seen = set()

    for group in groups:
        name_col = next((m[1] for m in group if m[0] == "name"), group[0][1])
        rank_col = next((m[1] for m in group if m[0] == "rank"), None)
        score_col = next((m[1] for m in group if m[0] == "score"), None)
        id_col = next((m[1] for m in group if m[0] == "student_id"), None)
        gender_col = next((m[1] for m in group if m[0] == "gender"), None)

        for row in rows[1:]:
            if name_col >= len(row) or not row[name_col]:
                continue
            name = str(row[name_col]).strip()
            if not name or len(name) < 2 or len(name) > 20:
                continue
            if name in seen:
                continue
            seen.add(name)

            rank = None
            total_score = 0.0

            if score_col is not None and score_col < len(row) and row[score_col] is not None:
                try:
                    total_score = float(str(row[score_col]))
                    rank = 0
                except (ValueError, TypeError):
                    pass

            if rank is None and rank_col is not None and rank_col < len(row) and row[rank_col] is not None:
                try:
                    rank = int(float(str(row[rank_col])))
                    total_score = float(1000000 - rank)
                except (ValueError, TypeError):
                    pass

            student_id = None
            if id_col is not None and id_col < len(row) and row[id_col] is not None:
                student_id = str(row[id_col]).strip()

            gender = None
            if gender_col is not None and gender_col < len(row) and row[gender_col] is not None:
                gv = str(row[gender_col]).strip()
                gender = _normalize_gender(gv)

            students.append({
                "name": name,
                "rank": rank,
                "student_id": student_id,
                "total_score": total_score,
                "gender": gender,
            })

    return students, None


@router.post("/import-xlsx")
def import_xlsx(
    file: UploadFile = File(...),
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    client_ip = _get_client_ip(request) if request else "unknown"
    file_upload_rate_limiter.check(f"upload:{client_ip}")

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式文件")

    if file.content_type and file.content_type not in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    ):
        raise HTTPException(status_code=400, detail="不支持的文件类型，请上传 .xlsx 格式文件")

    raw = file.file.read()
    _validate_file_magic(raw, b"PK\x03\x04", ".xlsx")

    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(raw) > max_bytes:
        raise HTTPException(
            status_code=400,
            detail=f"文件大小超过{settings.MAX_UPLOAD_SIZE_MB}MB限制"
        )

    students, error = _parse_xlsx_students(raw)

    if error:
        raise HTTPException(status_code=400, detail=error)

    if not students:
        raise HTTPException(status_code=400, detail="未解析到有效学生数据")

    students.sort(key=lambda s: (-s["total_score"], s["rank"] if s["rank"] is not None else 999999))

    existing = {s.name for s in db.query(models.Student).all()}
    added = 0
    updated = 0

    for idx, s in enumerate(students):
        db_student = db.query(models.Student).filter(models.Student.name == s["name"]).first()
        if db_student:
            db_student.student_id = s["student_id"]
            db_student.rank = s["rank"]
            db_student.total_score = s["total_score"]
            if s.get("gender") is not None:
                db_student.gender = s["gender"]
            updated += 1
        else:
            db.add(models.Student(
                name=s["name"],
                student_id=s["student_id"],
                rank=s["rank"],
                total_score=s["total_score"],
                gender=s.get("gender"),
            ))
            added += 1

    db.commit()

    if db.query(models.Seat).count() == 0:
        state = get_seat_state(db)
        _rebuild_seats(db, state.groups_count, state.rows_per_group, state.cols_per_group, state.has_pairs)
        db.commit()

    log_action(
        action="import_xlsx",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"filename": file.filename, "added": added, "updated": updated},
        ip=_get_client_ip(request) if request else None,
    )

    return {
        "success": True,
        "message": f"成功导入 {added} 名新学生，更新 {updated} 名已有学生",
        "total": len(students),
        "added": added,
        "updated": updated,
        "students": students,
    }


@router.post("/clear-seats")
def clear_all_seats(
    req: schemas.DestructiveActionConfirm,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if not auth.verify_teacher_password(db, req.teacher_password):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="教师密码验证失败")

    db.query(models.Seat).delete()
    db.query(models.SeatState).delete()
    db.commit()

    log_action(
        action="clear_seats",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        ip=_get_client_ip(request) if request else None,
    )

    return {"success": True, "message": "所有座位数据已清空"}


@router.get("/layout")
def get_layout(db: Session = Depends(get_db)):
    state = get_seat_state(db)
    return {
        "groupsCount": state.groups_count,
        "rowsPerGroup": state.rows_per_group,
        "colsPerGroup": state.cols_per_group,
        "hasPairs": state.has_pairs,
        "totalSeats": state.groups_count * state.rows_per_group * state.cols_per_group + 2,
    }


def _rebuild_seats(db: Session, groups_count: int, rows_per_group: int, cols_per_group: int, has_pairs: bool):
    db.query(models.Seat).delete()
    platform_seats = [
        models.Seat(id=1, label="讲台左", seat_type="platformLeft", group_idx=None, row=None, col=None),
        models.Seat(id=2, label="讲台右", seat_type="platformRight", group_idx=None, row=None, col=None),
    ]
    for s in platform_seats:
        db.add(s)
    seat_id = 3
    for g in range(groups_count):
        for r in range(rows_per_group):
            for c in range(cols_per_group):
                db.add(models.Seat(
                    id=seat_id,
                    label=f"第{g+1}组第{r+1}排第{c+1}列",
                    seat_type="group",
                    group_idx=g + 1,
                    row=r + 1,
                    col=c + 1,
                ))
                seat_id += 1
    db.flush()


@router.post("/layout")
def update_layout(
    req: schemas.LayoutUpdate,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if req.groups_count < 1 or req.rows_per_group < 1 or req.cols_per_group < 2:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="组数、行数至少为1，列数至少为2")

    total_group_seats = req.groups_count * req.rows_per_group * req.cols_per_group
    student_count = db.query(models.Student).count()
    if student_count > 0 and total_group_seats < student_count:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"座位总数({total_group_seats})不能少于已导入的学生人数({student_count})"
        )

    state = get_seat_state(db)
    old_layout = {
        "groups": state.groups_count,
        "rows": state.rows_per_group,
        "cols": state.cols_per_group,
        "pairs": state.has_pairs,
    }
    state.groups_count = req.groups_count
    state.rows_per_group = req.rows_per_group
    state.cols_per_group = req.cols_per_group
    state.has_pairs = req.has_pairs
    state.current_index = 0

    students = db.query(models.Student).all()
    for student in students:
        student.assigned = False
        student.seat_id = None
        student.selection_method = None

    _rebuild_seats(db, req.groups_count, req.rows_per_group, req.cols_per_group, req.has_pairs)
    db.commit()

    log_action(
        action="update_layout",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"old": old_layout, "new": req.model_dump()},
        ip=_get_client_ip(request) if request else None,
    )

    return {"success": True, "message": f"教室布局已更新：{req.groups_count}组 × {req.rows_per_group}排 × {req.cols_per_group}列"}


@router.post("/seats/{seat_id}/toggle-disable")
def toggle_seat_disable(
    seat_id: int,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    seat = db.query(models.Seat).filter(models.Seat.id == seat_id).first()
    if not seat:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="座位不存在")

    seat.is_disabled = not seat.is_disabled
    if seat.is_disabled and seat.is_occupied:
        student = db.query(models.Student).filter(models.Student.name == seat.occupant_name).first()
        if student:
            student.assigned = False
            student.seat_id = None
            student.selection_method = None
            student.round_id = None
        seat.is_occupied = False
        seat.occupant_name = None
        seat.occupant_id = None

    db.commit()
    return {
        "success": True,
        "message": "座位已禁用" if seat.is_disabled else "座位已启用",
        "isDisabled": seat.is_disabled,
    }


@router.post("/undo-round")
def undo_round_assignments(
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    state = get_seat_state(db)
    state.current_index = 0

    seats = db.query(models.Seat).all()
    cleared_count = 0
    for seat in seats:
        if seat.is_occupied:
            seat.is_occupied = False
            seat.occupant_name = None
            seat.occupant_id = None
            cleared_count += 1

    students = db.query(models.Student).all()
    for student in students:
        if student.assigned:
            student.assigned = False
            student.seat_id = None
            student.selection_method = None
            student.round_id = None

    db.commit()
    return {"success": True, "message": f"已撤销本轮所有选座操作，共清除 {cleared_count} 个座位", "currentIndex": 0}


@router.post("/reset-data")
def reset_all_data(
    req: schemas.DestructiveActionConfirm,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if not auth.verify_teacher_password(db, req.teacher_password):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="教师密码验证失败")

    db.query(models.Wish).delete()
    db.query(models.Invitation).delete()
    db.query(models.SelectionRound).delete()
    db.query(models.Student).delete()
    db.query(models.Seat).delete()
    db.query(models.OperationHistory).delete()
    db.query(models.SeatState).delete()
    db.commit()

    log_action(
        action="reset_all_data",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        ip=_get_client_ip(request) if request else None,
    )

    return {"success": True, "message": "所有数据已重置，请重新导入学生和座位"}


@router.post("/auto-lock-wishes")
def auto_lock_wishes(
    req: schemas.AutoLockRequest,
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    if req.quota <= 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="配额必须大于0")

    students = db.query(models.Student).filter(models.Student.assigned == False).all()
    students.sort(key=lambda s: (-s.total_score, s.rank if s.rank is not None else 0, s.name))

    if len(students) > 1:
        i = 0
        while i < len(students):
            j = i
            while j < len(students) and students[j].total_score == students[i].total_score:
                j += 1
            if j - i > 1:
                random.shuffle(students[i:j])
            i = j

    top_students = students[:req.quota]
    locked = []
    skipped = []

    for student in top_students:
        wishes = db.query(models.Wish).filter(
            models.Wish.student_name == student.name
        ).order_by(models.Wish.priority).all()

        assigned = False
        for wish in wishes:
            seat = db.query(models.Seat).filter(models.Seat.id == wish.seat_id).first()
            if seat and not seat.is_occupied and not seat.is_disabled:
                seat.is_occupied = True
                seat.occupant_name = student.name
                seat.occupant_id = student.id
                student.assigned = True
                student.seat_id = seat.id
                locked.append({"name": student.name, "seat": seat.label, "priority": wish.priority})
                assigned = True
                break

        if not assigned:
            skipped.append(student.name)

    state = get_seat_state(db)
    sorted_students = sorted(students, key=lambda s: (-s.total_score, s.rank if s.rank is not None else 0, s.name))
    state.current_index = 0
    for i, s in enumerate(sorted_students):
        if s.assigned:
            state.current_index = i + 1
        else:
            break

    db.commit()

    log_action(
        action="auto_lock_wishes",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"quota": req.quota, "locked_count": len(locked), "skipped_count": len(skipped)},
        ip=_get_client_ip(request) if request else None,
    )

    return {
        "success": True,
        "locked": locked,
        "skipped": skipped,
        "message": f"已锁定 {len(locked)} 名学生，{len(skipped)} 名志愿全部冲突需要重选",
    }


@router.post("/seed-demo")
def seed_demo_data(
    request: Request = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(auth.require_role("teacher"))
):
    existing = db.query(models.Student).first()
    if existing:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="已有数据，请先重置")

    demo_students = [
        {"id": 1, "name": "张伟", "totalScore": 685, "gender": "m"},
        {"id": 2, "name": "李娜", "totalScore": 678, "gender": "f"},
        {"id": 3, "name": "王芳", "totalScore": 672, "gender": "f"},
        {"id": 4, "name": "刘洋", "totalScore": 665, "gender": "m"},
        {"id": 5, "name": "陈静", "totalScore": 660, "gender": "f"},
        {"id": 6, "name": "杨光", "totalScore": 655, "gender": "m"},
        {"id": 7, "name": "赵敏", "totalScore": 650, "gender": "f"},
        {"id": 8, "name": "黄磊", "totalScore": 645, "gender": "m"},
        {"id": 9, "name": "周杰", "totalScore": 640, "gender": "m"},
        {"id": 10, "name": "吴桐", "totalScore": 635, "gender": "m"},
        {"id": 11, "name": "徐慧", "totalScore": 630, "gender": "f"},
        {"id": 12, "name": "孙鹏", "totalScore": 625, "gender": "m"},
        {"id": 13, "name": "马丽", "totalScore": 620, "gender": "f"},
        {"id": 14, "name": "朱峰", "totalScore": 615, "gender": "m"},
        {"id": 15, "name": "胡洁", "totalScore": 610, "gender": "f"},
        {"id": 16, "name": "郭瑞", "totalScore": 605, "gender": "m"},
        {"id": 17, "name": "林霞", "totalScore": 600, "gender": "f"},
        {"id": 18, "name": "何勇", "totalScore": 595, "gender": "m"},
        {"id": 19, "name": "高远", "totalScore": 590, "gender": "m"},
        {"id": 20, "name": "罗琳", "totalScore": 585, "gender": "f"},
        {"id": 21, "name": "梁宇", "totalScore": 580, "gender": "m"},
        {"id": 22, "name": "宋雨", "totalScore": 575, "gender": "f"},
        {"id": 23, "name": "唐明", "totalScore": 570, "gender": "m"},
        {"id": 24, "name": "韩雪", "totalScore": 565, "gender": "f"},
        {"id": 25, "name": "曹阳", "totalScore": 560, "gender": "m"},
        {"id": 26, "name": "邓超", "totalScore": 555, "gender": "m"},
        {"id": 27, "name": "许晴", "totalScore": 550, "gender": "f"},
        {"id": 28, "name": "彭涛", "totalScore": 545, "gender": "m"},
        {"id": 29, "name": "沈逸", "totalScore": 540, "gender": "m"},
        {"id": 30, "name": "曾辉", "totalScore": 535, "gender": "m"},
        {"id": 31, "name": "卢芳", "totalScore": 530, "gender": "f"},
        {"id": 32, "name": "蒋欣", "totalScore": 525, "gender": "f"},
        {"id": 33, "name": "蔡康", "totalScore": 520, "gender": "m"},
        {"id": 34, "name": "贾玲", "totalScore": 515, "gender": "f"},
        {"id": 35, "name": "丁伟", "totalScore": 500, "gender": "m"},
        {"id": 36, "name": "魏芳", "totalScore": 495, "gender": "f"},
        {"id": 37, "name": "薛之", "totalScore": 490, "gender": "m"},
        {"id": 38, "name": "叶舒", "totalScore": 485, "gender": "f"},
        {"id": 39, "name": "阎宁", "totalScore": 480, "gender": "f"},
        {"id": 40, "name": "余晨", "totalScore": 475, "gender": "m"},
        {"id": 41, "name": "潘越", "totalScore": 470, "gender": "m"},
        {"id": 42, "name": "戴玉", "totalScore": 465, "gender": "f"},
        {"id": 43, "name": "夏冰", "totalScore": 460, "gender": "f"},
        {"id": 44, "name": "田蜜", "totalScore": 455, "gender": "f"},
        {"id": 45, "name": "任航", "totalScore": 450, "gender": "m"},
        {"id": 46, "name": "姜文", "totalScore": 445, "gender": "m"},
        {"id": 47, "name": "范凯", "totalScore": 440, "gender": "m"},
        {"id": 48, "name": "方圆", "totalScore": 435, "gender": "f"},
        {"id": 49, "name": "石磊", "totalScore": 430, "gender": "m"},
        {"id": 50, "name": "姚远", "totalScore": 425, "gender": "m"},
        {"id": 51, "name": "谭笑", "totalScore": 420, "gender": "f"},
        {"id": 52, "name": "廖静", "totalScore": 415, "gender": "f"},
        {"id": 53, "name": "邹凯", "totalScore": 410, "gender": "m"},
    ]

    for s in demo_students:
        student = models.Student(**s)
        db.add(student)

    demo_seats = [
        {"id": 1, "label": "讲台左", "seat_type": "platformLeft", "group_idx": None, "row": None, "col": None},
        {"id": 2, "label": "讲台右", "seat_type": "platformRight", "group_idx": None, "row": None, "col": None},
    ]
    state = get_seat_state(db)
    g_cnt = state.groups_count
    r_cnt = state.rows_per_group
    c_cnt = state.cols_per_group
    seat_id = 3
    for g in range(g_cnt):
        for r in range(r_cnt):
            for c in range(c_cnt):
                demo_seats.append({
                    "id": seat_id,
                    "label": f"第{g+1}组第{r+1}排第{c+1}列",
                    "seat_type": "group",
                    "group_idx": g + 1,
                    "row": r + 1,
                    "col": c + 1,
                })
                seat_id += 1

    for s in demo_seats:
        seat = models.Seat(**s)
        db.add(seat)

    state = models.SeatState()
    db.add(state)

    db.commit()

    log_action(
        action="seed_demo",
        user_role=current_user["role"],
        user_identity=current_user.get("sub", "unknown"),
        detail={"students_count": len(demo_students), "seats_count": len(demo_seats)},
        ip=_get_client_ip(request) if request else None,
    )

    return {"success": True, "message": f"已添加{len(demo_students)}名学生和{len(demo_seats)}个座位"}